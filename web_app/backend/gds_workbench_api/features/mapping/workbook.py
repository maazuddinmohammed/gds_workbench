"""The agreed Mapping layout: one System workbook, one sheet per target Entity.

Inputs are a projection of saved records. Authorization, record selection and
template interpretation belong to the read service, not the XLSX renderer.
"""

import re
from dataclasses import dataclass
from io import BytesIO
from math import ceil
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

MAPPING_ATTRIBUTE_COLUMNS = (
    "Index",
    "TargetColumn",
    "DataType",
    "NotNull",
    "IsNaturalKey",
    "IsSurrogateKey",
    "TransformationLogic",
    "DefaultRecord",
    "SourceColumns",
)
MAX_MAPPING_SHEETS = 200
MAX_MAPPING_ROWS_PER_SHEET = 10_000
MAX_MAPPING_ROWS = 50_000
MAX_MAPPING_CELL_CHARACTERS = 32_767
MAX_MAPPING_TEXT_CHARACTERS = 16 * 1024 * 1024
MAX_MAPPING_XLSX_BYTES = 32 * 1024 * 1024
MAX_MAPPING_UNCOMPRESSED_BYTES = 96 * 1024 * 1024

_INVALID_XML_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_INVALID_SHEET_TITLE = re.compile(r"[\[\]:*?/\\\x00-\x1f\x7f]")


class MappingWorkbookBuildError(ValueError):
    """A bounded export failure whose message never contains record content."""


@dataclass(frozen=True, slots=True)
class MappingWorkbookAttribute:
    ordinal_position: int
    target_column: str
    data_type: str
    is_nullable: bool
    is_natural_key: bool
    is_surrogate_key: bool
    transformation_logic: str | None
    default_record: str | None
    # Each item is (schema, object name, Attribute name), for either source kind.
    source_columns: tuple[tuple[str, str, str], ...]


@dataclass(frozen=True, slots=True)
class MappingWorkbookEntity:
    target_layer: str
    target_schema: str
    target_table_name: str
    target_table_description: str | None
    entity_type: str
    dependency_order: int | None
    # Each item is (schema, object name), for either physical or modeled sources.
    source_tables: tuple[tuple[str, str], ...]
    filter_criteria: str | None
    sample_query: str | None
    attributes: tuple[MappingWorkbookAttribute, ...]


def build_mapping_workbook(
    *,
    model_name: str,
    tenant_name: str,
    system_code: str,
    entities: tuple[MappingWorkbookEntity, ...],
) -> bytes:
    """Render a single System's projected records without filesystem or DB access."""
    if not entities:
        raise MappingWorkbookBuildError("No saved Mapping Entities are available to export.")
    if len(entities) > MAX_MAPPING_SHEETS:
        raise MappingWorkbookBuildError("Mapping workbook exceeds its sheet limit.")

    workbook = Workbook()
    initial_sheet = workbook.active
    if initial_sheet is None:
        raise MappingWorkbookBuildError("Mapping workbook initialization failed.")
    workbook.remove(initial_sheet)
    workbook.properties.creator = "Atlas"
    body_font = Font(name="Arial", size=10)
    header_font = Font(name="Arial", size=10, bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor="EAE5DF")
    border_side = Side(style="thin", color="808080")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    widths = (26, 38, 18, 12, 17, 18, 60, 26, 55)
    used_titles: set[str] = set()
    text_characters = 0
    total_rows = 0

    try:
        for entity in sorted(
            entities,
            key=lambda item: (
                item.dependency_order is None,
                item.dependency_order or 0,
                item.target_schema,
                item.target_table_name,
                item.target_layer,
            ),
        ):
            # This block is the complete, ordered top-section field mapping.
            header_rows = (
                ("ModelName", model_name),
                ("TenantName", tenant_name),
                ("SystemCode", system_code),
                ("TargetLayer", entity.target_layer),
                ("TargetSchema", entity.target_schema),
                ("TargetTableName", entity.target_table_name),
                ("TargetTableDescription", entity.target_table_description),
                ("EntityType", entity.entity_type),
                ("DependencyOrder", entity.dependency_order),
                ("SourceTables", "\n".join("|".join(source) for source in entity.source_tables)),
                ("FilterCriteria", entity.filter_criteria),
                ("SampleQuery", entity.sample_query),
            )
            grid_header_row = len(header_rows) + 2  # Exactly one blank separating row.
            sheet_rows = grid_header_row + len(entity.attributes)
            total_rows += sheet_rows
            if sheet_rows > MAX_MAPPING_ROWS_PER_SHEET or total_rows > MAX_MAPPING_ROWS:
                raise MappingWorkbookBuildError("Mapping workbook exceeds its row limit.")

            # Full names remain in the header; tab titles obey Excel's restrictions.
            base_title = _INVALID_SHEET_TITLE.sub("_", entity.target_table_name).strip(" '")
            base_title = base_title or "Entity"
            title = base_title[:31].rstrip(" '")
            suffix_number = 2
            while title.casefold() in used_titles or title.casefold() == "history":
                suffix = f"_{suffix_number}"
                title = base_title[: 31 - len(suffix)].rstrip(" '") + suffix
                suffix_number += 1
            used_titles.add(title.casefold())
            worksheet = workbook.create_sheet(title)

            attribute_rows = [
                (
                    attribute.ordinal_position,
                    attribute.target_column,
                    attribute.data_type,
                    not attribute.is_nullable,
                    attribute.is_natural_key,
                    attribute.is_surrogate_key,
                    attribute.transformation_logic,
                    attribute.default_record,
                    "\n".join("|".join(source) for source in attribute.source_columns),
                )
                for attribute in sorted(entity.attributes, key=lambda item: item.ordinal_position)
            ]
            for row_number, values in enumerate(
                [*header_rows, (), MAPPING_ATTRIBUTE_COLUMNS, *attribute_rows], start=1
            ):
                # Validate before assignment: openpyxl otherwise truncates long strings.
                for value in values:
                    if isinstance(value, str):
                        if _INVALID_XML_TEXT.search(value):
                            raise MappingWorkbookBuildError(
                                "Mapping workbook contains text Excel cannot represent."
                            )
                        # Excel's string limit counts UTF-16 units, including emoji pairs.
                        length = len(value.encode("utf-16-le")) // 2
                        if length > MAX_MAPPING_CELL_CHARACTERS:
                            raise MappingWorkbookBuildError(
                                "Mapping workbook contains a cell exceeding Excel's text limit."
                            )
                        text_characters += length
                        if text_characters > MAX_MAPPING_TEXT_CHARACTERS:
                            raise MappingWorkbookBuildError(
                                "Mapping workbook exceeds its text limit."
                            )
                    elif value is not None and type(value) not in (int, bool):
                        raise MappingWorkbookBuildError(
                            "Mapping workbook contains an unsupported cell value."
                        )
                worksheet.append(values)
                line_count = 1
                for column_number, cell in enumerate(worksheet[row_number], start=1):
                    cell.font = body_font
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if (row_number <= len(header_rows) and column_number <= 2) or (
                        row_number >= grid_header_row
                    ):
                        cell.border = cell_border
                    if isinstance(cell.value, str):
                        # Mapping SQL/text beginning '=' or '#N/A' is literal documentation.
                        cell.data_type = "s"
                        line_count = max(
                            line_count,
                            sum(
                                max(1, ceil(len(line) / widths[column_number - 1]))
                                for line in cell.value.split("\n")
                            ),
                        )
                    elif isinstance(cell.value, bool):
                        cell.alignment = Alignment(horizontal="center", vertical="top")
                    elif isinstance(cell.value, int):
                        cell.number_format = "0"
                    if row_number == grid_header_row or (
                        row_number <= len(header_rows) and column_number == 1
                    ):
                        cell.font = header_font
                        cell.fill = header_fill
                worksheet.row_dimensions[row_number].height = min(409, max(18, line_count * 15))

            for column_number, width in enumerate(widths, start=1):
                worksheet.column_dimensions[chr(64 + column_number)].width = width
            worksheet.auto_filter.ref = f"A{grid_header_row}:I{sheet_rows}"
            worksheet.sheet_view.showGridLines = False
            worksheet.sheet_properties.pageSetUpPr.fitToPage = True
            worksheet.page_setup.orientation = "landscape"
            worksheet.page_setup.fitToWidth = 1
            worksheet.page_setup.fitToHeight = 0
            worksheet.print_title_rows = f"{grid_header_row}:{grid_header_row}"

        output = BytesIO()
        workbook.save(output)
        content = output.getvalue()
        if len(content) > MAX_MAPPING_XLSX_BYTES:
            raise MappingWorkbookBuildError("Mapping workbook exceeds its download size limit.")
        with ZipFile(BytesIO(content)) as package:
            if (
                sum(member.file_size for member in package.infolist())
                > MAX_MAPPING_UNCOMPRESSED_BYTES
            ):
                raise MappingWorkbookBuildError("Mapping workbook exceeds its package size limit.")
        return content
    finally:
        workbook.close()
