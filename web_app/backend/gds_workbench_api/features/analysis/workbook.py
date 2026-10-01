"""Single-sheet Analysis export matching the results table's data columns.

The read service supplies current saved findings. Selection, authorization and
download filenames remain separate from this in-memory renderer.
"""

import re
from io import BytesIO
from math import ceil
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .read_contracts import AnalysisFindingSummary

ANALYSIS_WORKBOOK_COLUMNS = (
    "FromObjectSchema",
    "FromObjectName",
    "FromAttributeName",
    "ToObjectSchema",
    "ToObjectName",
    "ToAttributeName",
    "Relationship",
    "InferredCardinality",
    "ObservedCardinality",
    "Confidence",
    "FromRowCount",
    "ToRowCount",
    "SourceMissingTargetCount",
    "UnusedTargetCount",
    "SourceMissingTargetPercent",
    "UnusedTargetPercent",
    "Status",
    "IsLocked",
)
MAX_ANALYSIS_FINDINGS = 50_000
MAX_ANALYSIS_CELL_CHARACTERS = 32_767
MAX_ANALYSIS_TEXT_CHARACTERS = 16 * 1024 * 1024
MAX_ANALYSIS_XLSX_BYTES = 32 * 1024 * 1024
MAX_ANALYSIS_UNCOMPRESSED_BYTES = 96 * 1024 * 1024
_INVALID_XML_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


class AnalysisWorkbookBuildError(ValueError):
    """A safe export failure with no finding content in its message."""


def build_analysis_workbook(*, findings: tuple[AnalysisFindingSummary, ...]) -> bytes:
    """Render one bordered Analysis sheet, preserving the supplied finding order."""
    if len(findings) > MAX_ANALYSIS_FINDINGS:
        raise AnalysisWorkbookBuildError("Analysis workbook exceeds its row limit.")

    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:
        raise AnalysisWorkbookBuildError("Analysis workbook initialization failed.")
    worksheet.title = "Analysis"
    workbook.properties.creator = "Atlas"
    worksheet.append(ANALYSIS_WORKBOOK_COLUMNS)
    text_characters = sum(len(label) for label in ANALYSIS_WORKBOOK_COLUMNS)
    widths = (24, 28, 28, 24, 28, 28, 24, 24, 30, 14, 20, 20, 26, 24, 28, 24, 16, 12)
    border_side = Side(style="thin", color="808080")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    body_font = Font(name="Arial", size=10)
    header_font = Font(name="Arial", size=10, bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor="EAE5DF")

    try:
        for finding in findings:
            values = list(_analysis_row(finding))
            for index, value in enumerate(values):
                if type(value) is int and abs(value) >= 10**15:
                    # Excel numbers have 15 significant digits; preserve larger counts as text.
                    value = str(value)
                    values[index] = value
                if isinstance(value, str):
                    if _INVALID_XML_TEXT.search(value):
                        raise AnalysisWorkbookBuildError(
                            "Analysis workbook contains text Excel cannot represent."
                        )
                    length = len(value.encode("utf-16-le")) // 2
                    if length > MAX_ANALYSIS_CELL_CHARACTERS:
                        raise AnalysisWorkbookBuildError(
                            "Analysis workbook contains a cell exceeding Excel's text limit."
                        )
                    text_characters += length
                    if text_characters > MAX_ANALYSIS_TEXT_CHARACTERS:
                        raise AnalysisWorkbookBuildError(
                            "Analysis workbook exceeds its text limit."
                        )
            # Validate before assignment so openpyxl cannot silently truncate text.
            worksheet.append(values)

        for row_number, row in enumerate(worksheet.iter_rows(), start=1):
            line_count = 1
            for column_number, cell in enumerate(row, start=1):
                cell.font = header_font if row_number == 1 else body_font
                cell.border = cell_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_number == 1:
                    cell.fill = header_fill
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # SQL-like names and Excel errors are literal text.
                    line_count = max(
                        line_count,
                        sum(
                            max(1, ceil(len(line) / widths[column_number - 1]))
                            for line in cell.value.split("\n")
                        ),
                    )
                elif isinstance(cell.value, bool):
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00%" if column_number in (15, 16) else "#,##0"
            worksheet.row_dimensions[row_number].height = min(409, max(18, line_count * 15))

        for column_number, width in enumerate(widths, start=1):
            worksheet.column_dimensions[chr(64 + column_number)].width = width
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:R{len(findings) + 1}"
        worksheet.sheet_view.showGridLines = False
        if worksheet.sheet_properties.pageSetUpPr is not None:
            worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.print_title_rows = "1:1"

        output = BytesIO()
        workbook.save(output)
        content = output.getvalue()
        if len(content) > MAX_ANALYSIS_XLSX_BYTES:
            raise AnalysisWorkbookBuildError("Analysis workbook exceeds its download size limit.")
        with ZipFile(BytesIO(content)) as package:
            if (
                sum(member.file_size for member in package.infolist())
                > MAX_ANALYSIS_UNCOMPRESSED_BYTES
            ):
                raise AnalysisWorkbookBuildError(
                    "Analysis workbook exceeds its package size limit."
                )
        return content
    finally:
        workbook.close()


def _analysis_row(finding: AnalysisFindingSummary) -> tuple[str | int | float | bool, ...]:
    """Project the visible columns, including their missing-evidence states."""
    unavailable = "Not validated" if finding.validation_state == "unvalidated" else "Unavailable"
    labels = [
        finding.relationship_kind.replace("_", " "),
        finding.inferred_cardinality.replace("_", " "),
        finding.observed_cardinality.replace("_", " ")
        if finding.observed_cardinality
        else unavailable,
    ]
    labels = [label[:1].upper() + label[1:] for label in labels]
    if finding.cardinality_mismatch:
        labels[2] += "\nDiffers from inference"
    counts = (finding.source_missing_target_count, finding.unused_target_count)
    percentages = tuple(
        percent / 100
        if percent is not None
        else (unavailable if finding.validation_state == "unvalidated" or count is None else "N/A")
        for percent, count in zip(
            (finding.source_missing_target_percent, finding.unused_target_percent),
            counts,
            strict=True,
        )
    )
    # This tuple follows ANALYSIS_WORKBOOK_COLUMNS exactly. No checkbox, IDs or audit payloads.
    return (
        finding.from_endpoint.object_schema,
        finding.from_endpoint.object_name,
        finding.from_endpoint.attribute_name,
        finding.to_endpoint.object_schema,
        finding.to_endpoint.object_name,
        finding.to_endpoint.attribute_name,
        *labels,
        finding.relationship_confidence,
        finding.from_row_count if finding.from_row_count is not None else "Not profiled",
        finding.to_row_count if finding.to_row_count is not None else "Not profiled",
        *(count if count is not None else unavailable for count in counts),
        *percentages,
        finding.status[:1].upper() + finding.status[1:],
        finding.is_locked,
    )
