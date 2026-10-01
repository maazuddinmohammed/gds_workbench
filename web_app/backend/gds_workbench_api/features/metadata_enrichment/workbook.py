"""Two-sheet Model enrichment dictionary, with explicit export column mappings."""

import re
from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gds_workbench_api.features.model_input_scope.contracts import (
    ModelInputScopeAttribute,
    ModelInputScopeObject,
)

DATA_DICTIONARY_COLUMNS = (
    "ObjectSchemaName",
    "ObjectTableName",
    "ObjectDescription",
    "AttributeName",
    "AttributeDescription",
    "InferredDataType",
    "MetadataDataType",
    "OrdinalPosition",
    "IsNaturalKey",
    "IsPrimaryKey",
    "IsNullable",
    "IsPII",
)
TECHNICAL_DICTIONARY_COLUMNS = (
    "ObjectSchemaName",
    "ObjectTableName",
    "AttributeName",
    "OrdinalPosition",
    "AttributeDataType",
    "TotalRows",
    "NotNullRows",
    "NullRows",
    "BlankRows",
    "DistinctValues",
    "PercentPopulated",
    "PercentNull",
    "PercentBlank",
    "PercentDistinct",
    "PercentDuplicates",
    "MinDataLength",
    "MaxDataLength",
    "AvgDataLength",
    "ProfiledAt",
    "RowScope",
    "BatchId",
)
PROFILE_FIELDS = (
    "row_count",
    "non_null_count",
    "null_count",
    "blank_count",
    "distinct_count",
    "percent_populated",
    "percent_null",
    "percent_blank",
    "percent_distinct",
    "percent_duplicates",
    "min_data_length",
    "max_data_length",
    "avg_data_length",
    "updated_time",
    "row_scope",
    "batch_id",
)
MAX_DICTIONARY_ROWS = 50_000


def build_enrichment_workbook(
    rows: Sequence[tuple[ModelInputScopeObject, ModelInputScopeAttribute | None]],
) -> bytes:
    if len(rows) > MAX_DICTIONARY_ROWS:
        raise ValueError("Enrichment export exceeds its row limit.")
    workbook = Workbook()
    data = workbook.active
    if data is None:
        raise ValueError("Workbook initialization failed.")
    data.title = "Data Dictionary"
    technical = workbook.create_sheet("Technical Data Dictionary")
    data.append(DATA_DICTIONARY_COLUMNS)
    technical.append(TECHNICAL_DICTIONARY_COLUMNS)
    side = Side(style="thin", color="808080")
    border = Border(left=side, right=side, top=side, bottom=side)
    text_size = 0
    try:
        for obj, attr in rows:
            findings = attr.enrichment if attr else None
            data_values = [
                obj.object_schema,
                obj.object_name,
                obj.object_description,
                attr.attribute_name if attr else None,
                attr.attribute_description if attr else None,
                attr.attribute_inferred_data_type if attr else None,
                attr.attribute_data_type if attr else None,
                attr.attribute_ordinal_position if attr else None,
                *[
                    getattr(findings, key) if findings else None
                    for key in ("is_natural_key", "is_primary_key", "is_nullable", "is_pii")
                ],
            ]
            profile = attr.profile if attr and attr.profile else {}
            technical_values = [
                obj.object_schema,
                obj.object_name,
                attr.attribute_name if attr else None,
                attr.attribute_ordinal_position if attr else None,
                attr.attribute_data_type if attr else None,
                *[profile.get(key) for key in PROFILE_FIELDS],
            ]
            # Validate before openpyxl can truncate strings or echo invalid content.
            for sheet, values in ((data, data_values), (technical, technical_values)):
                normalized: list[str | int | float | bool | None] = []
                for value in values:
                    if isinstance(value, (date, datetime)):
                        value = value.isoformat()
                    if type(value) is int and abs(value) >= 10**15:
                        value = str(value)
                    if isinstance(value, Decimal):
                        value = float(value)
                    if isinstance(value, str):
                        if (
                            re.search(
                                r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", value
                            )
                            or len(value.encode("utf-16-le")) // 2 > 32767
                        ):
                            raise ValueError("Enrichment export contains unsupported cell text.")
                        text_size += len(value.encode("utf-8"))
                        if text_size > 16 * 1024 * 1024:
                            raise ValueError("Enrichment export exceeds its text limit.")
                    elif value is not None and not isinstance(value, (bool, int, float)):
                        raise ValueError("Enrichment export contains an unsupported value.")
                    normalized.append(value)
                sheet.append(normalized)
        for sheet in (data, technical):
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cells in sheet:
                for cell in cells:
                    if not isinstance(cell, Cell):
                        continue
                    if isinstance(cell.value, str):
                        # Keep metadata text literal, including =, +, -, and @ prefixes.
                        cell.data_type = "s"
                    cell.border = border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.font = Font(name="Arial", size=10, bold=cell.row == 1)
                    if cell.row == 1:
                        cell.fill = PatternFill(fill_type="solid", fgColor="EAE5DF")
            for index, cell in enumerate(sheet[1], 1):
                sheet.column_dimensions[get_column_letter(index)].width = (
                    52 if "Description" in str(cell.value) else 24
                )
        output = BytesIO()
        workbook.save(output)
        content = output.getvalue()
        with ZipFile(BytesIO(content)) as archive:
            if (
                len(content) > 32 * 1024 * 1024
                or sum(info.file_size for info in archive.infolist()) > 96 * 1024 * 1024
            ):
                raise ValueError("Enrichment workbook exceeds its size limit.")
        return content
    finally:
        workbook.close()
