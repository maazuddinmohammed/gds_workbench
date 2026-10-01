from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

import pytest
from gds_workbench_api.features.mapping import workbook as mapping_workbook
from gds_workbench_api.features.mapping.workbook import (
    MappingWorkbookAttribute,
    MappingWorkbookBuildError,
    MappingWorkbookEntity,
    build_mapping_workbook,
)
from openpyxl import load_workbook


@pytest.fixture
def entity() -> MappingWorkbookEntity:
    return MappingWorkbookEntity(
        target_layer="Silver",
        target_schema="silver_crm",
        target_table_name="Customer",
        target_table_description="Current customers",
        entity_type="master",
        dependency_order=7,
        source_tables=(("bronze_crm", "Customer"), ("bronze_crm", "Address")),
        filter_criteria="c.IsActive = 1",
        sample_query="SELECT c.*\nFROM bronze_crm.Customer c",
        attributes=(
            MappingWorkbookAttribute(
                ordinal_position=4,
                target_column="FullName",
                data_type="string",
                is_nullable=True,
                is_natural_key=False,
                is_surrogate_key=False,
                transformation_logic="CONCAT(FirstName, ' ', LastName)",
                default_record=None,
                source_columns=(
                    ("bronze_crm", "Customer", "FirstName"),
                    ("bronze_crm", "Customer", "LastName"),
                ),
            ),
            MappingWorkbookAttribute(
                ordinal_position=1,
                target_column="CustomerId",
                data_type="int",
                is_nullable=False,
                is_natural_key=True,
                is_surrogate_key=False,
                transformation_logic=None,
                default_record=None,
                source_columns=(),
            ),
        ),
    )


def test_system_workbook_has_only_entity_sheets_and_the_agreed_two_sections(
    entity: MappingWorkbookEntity,
) -> None:
    dimension = replace(
        entity,
        target_layer="Gold",
        target_schema="gold_crm",
        target_table_name="DimCustomer",
        entity_type="dimension",
        dependency_order=0,
        attributes=(replace(entity.attributes[1], is_natural_key=False, is_surrogate_key=True),),
    )
    content = build_mapping_workbook(
        model_name="Customer model",
        tenant_name="Owning Tenant",
        system_code="CRM",
        entities=(entity, dimension),
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    try:
        assert workbook.sheetnames == ["DimCustomer", "Customer"]
        assert all(sheet.sheet_state == "visible" for sheet in workbook)
        customer = workbook["Customer"]
        assert list(customer.iter_rows(min_row=1, max_row=12, max_col=2, values_only=True)) == [
            ("ModelName", "Customer model"),
            ("TenantName", "Owning Tenant"),
            ("SystemCode", "CRM"),
            ("TargetLayer", "Silver"),
            ("TargetSchema", "silver_crm"),
            ("TargetTableName", "Customer"),
            ("TargetTableDescription", "Current customers"),
            ("EntityType", "master"),
            ("DependencyOrder", 7),
            ("SourceTables", "bronze_crm|Customer\nbronze_crm|Address"),
            ("FilterCriteria", "c.IsActive = 1"),
            ("SampleQuery", "SELECT c.*\nFROM bronze_crm.Customer c"),
        ]
        assert all(cell.value is None for cell in customer[13])
        assert tuple(cell.value for cell in customer[14]) == (
            "Index",
            "TargetColumn",
            "DataType",
            "NotNull",
            "IsNaturalKey",
            "IsSurrogateKey",
            "TransformationLogic",
            "DefaultRecord",
            "SourceColumns",
        )
        # Preserve saved ordinals, the unmapped Attribute, and native boolean key flags.
        assert tuple(cell.value for cell in customer[15]) == (
            1,
            "CustomerId",
            "int",
            True,
            True,
            False,
            None,
            None,
            None,
        )
        assert tuple(cell.value for cell in customer[16]) == (
            4,
            "FullName",
            "string",
            False,
            False,
            False,
            "CONCAT(FirstName, ' ', LastName)",
            None,
            "bronze_crm|Customer|FirstName\nbronze_crm|Customer|LastName",
        )
        assert customer.max_row == 16
        assert customer["I16"].alignment.wrap_text
        assert not customer.merged_cells
        assert tuple(cell.value for cell in workbook["DimCustomer"][15])[3:6] == (
            True,
            False,
            True,
        )
        assert workbook["DimCustomer"]["B9"].value == 0
    finally:
        workbook.close()


@pytest.mark.parametrize("literal", ["=1+1", '=HYPERLINK("https://example.test")', "#N/A"])
def test_mapping_text_is_preserved_as_literal_cells_without_formulas_or_external_links(
    entity: MappingWorkbookEntity,
    literal: str,
) -> None:
    entity = replace(
        entity,
        sample_query=literal,
        attributes=(
            replace(
                entity.attributes[0],
                transformation_logic=literal,
                default_record=literal,
            ),
        ),
    )
    content = build_mapping_workbook(
        model_name=literal,
        tenant_name=literal,
        system_code=literal,
        entities=(entity,),
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    try:
        for coordinate in ("B1", "B2", "B3", "B12", "G15", "H15"):
            cell = workbook["Customer"][coordinate]
            assert cell.value == literal
            assert cell.data_type == "s"
    finally:
        workbook.close()
    with ZipFile(BytesIO(content)) as package:
        assert not any("externallinks" in name.lower() for name in package.namelist())
        for name in package.namelist():
            payload = package.read(name)
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                assert b"<f" not in payload
            if name.endswith(".rels"):
                assert b'TargetMode="External"' not in payload


def test_sheet_titles_are_unique_valid_and_keep_full_names_in_headers(
    entity: MappingWorkbookEntity,
) -> None:
    names = ("Customer", "customer", "'A/B:C?D*E[F]G\\H'", "X" * 40, "X" * 39 + "Y", "History")
    entities = tuple(
        replace(entity, target_table_name=name, dependency_order=i) for i, name in enumerate(names)
    )
    workbook = load_workbook(
        BytesIO(
            build_mapping_workbook(
                model_name="Model",
                tenant_name="Tenant",
                system_code="CRM",
                entities=entities,
            )
        )
    )
    try:
        assert len(workbook.sheetnames) == len(names)
        assert workbook.sheetnames[:2] == ["Customer", "customer_2"]
        assert len({name.casefold() for name in workbook.sheetnames}) == len(names)
        assert all(1 <= len(name) <= 31 for name in workbook.sheetnames)
        assert all(
            not any(character in name for character in "[]:*?/\\") for name in workbook.sheetnames
        )
        assert all(name.casefold() != "history" for name in workbook.sheetnames)
        assert [sheet["B6"].value for sheet in workbook] == list(names)
    finally:
        workbook.close()


@pytest.mark.parametrize("text", ["x" * 32_768, "😀" * 16_384])
def test_oversized_cells_fail_without_truncation_or_content_in_the_error(
    entity: MappingWorkbookEntity,
    text: str,
) -> None:
    with pytest.raises(
        MappingWorkbookBuildError, match="cell exceeding Excel's text limit"
    ) as error:
        build_mapping_workbook(
            model_name="Model",
            tenant_name="Tenant",
            system_code="CRM",
            entities=(replace(entity, sample_query=text),),
        )
    assert text not in str(error.value)


@pytest.mark.parametrize(
    "text", ["private\x00text", "private\x0btext", "private\ufffftext", "\ud800"]
)
def test_unrepresentable_text_fails_without_echoing_the_value(
    entity: MappingWorkbookEntity,
    text: str,
) -> None:
    with pytest.raises(MappingWorkbookBuildError, match="text Excel cannot represent") as error:
        build_mapping_workbook(
            model_name="Model",
            tenant_name="Tenant",
            system_code="CRM",
            entities=(replace(entity, sample_query=text),),
        )
    assert text not in str(error.value)


def test_empty_export_does_not_create_a_placeholder_sheet() -> None:
    with pytest.raises(MappingWorkbookBuildError, match="No saved Mapping Entities"):
        build_mapping_workbook(
            model_name="Model",
            tenant_name="Tenant",
            system_code="CRM",
            entities=(),
        )


@pytest.mark.parametrize(
    "limit, message",
    [
        ("MAX_MAPPING_SHEETS", "sheet limit"),
        ("MAX_MAPPING_ROWS_PER_SHEET", "row limit"),
        ("MAX_MAPPING_ROWS", "row limit"),
        ("MAX_MAPPING_TEXT_CHARACTERS", "text limit"),
        ("MAX_MAPPING_XLSX_BYTES", "download size limit"),
        ("MAX_MAPPING_UNCOMPRESSED_BYTES", "package size limit"),
    ],
)
def test_export_bounds_fail_explicitly_instead_of_producing_partial_workbooks(
    entity: MappingWorkbookEntity,
    monkeypatch: pytest.MonkeyPatch,
    limit: str,
    message: str,
) -> None:
    monkeypatch.setattr(mapping_workbook, limit, 0)
    with pytest.raises(MappingWorkbookBuildError, match=message):
        build_mapping_workbook(
            model_name="Model",
            tenant_name="Tenant",
            system_code="CRM",
            entities=(entity,),
        )
