"""Missing Mapping rows are projected from modeled Attributes, never persisted by reads."""

# Reuse synthetic disposable fixture builders only.
# pyright: reportPrivateUsage=false
from io import BytesIO
from typing import Literal

from openpyxl import load_workbook

import pytest
from gds_etl_workbench.application.authorization import AuthorizationService
from gds_etl_workbench.domain.authorization import ActorKind, RequestPrincipal
from gds_workbench_api.database import WebPostgresDatabase
from gds_workbench_api.features.mapping.read_contracts import (
    MappingAttributeFilters,
    MappingFilters,
)
from gds_workbench_api.features.mapping.read_service import DatabaseMappingReviewService
from psycopg import sql

from tests.mcp.conftest import DisposablePostgres
from tests.mcp.database_test_support import require_row
from tests.web_backend.test_database_mapping_source_context import _seed_mapping_scope


@pytest.mark.parametrize("layer", ["logical", "dimensional"])
async def test_partial_mapping_reads_preserve_missing_attributes_and_empty_pair_history(
    web_postgres_database: DisposablePostgres,
    layer: Literal["logical", "dimensional"],
) -> None:
    scope = _seed_mapping_scope(
        web_postgres_database, dimensional=layer == "dimensional"
    )
    model_id, entity_id = scope.plan.model_id, scope.plan.pair.modeled_entity_id
    entity_type = scope.plan.modeled_entity_type
    with web_postgres_database.connect_owner() as connection:
        actor = require_row(
            connection.execute(
                "SELECT entra_tenant_id, entra_object_id FROM security.entra_principal_identity "
                "WHERE principal_id=%s",
                (scope.plan.actor_principal_id,),
            ).fetchone()
        )
        mapping_id = require_row(
            connection.execute(
                sql.SQL("""
            INSERT INTO workflow.mapping_object (
                model_id, modeled_entity_type, {}, source_system_id, mapping_transformation_document
            ) VALUES (%s, %s, %s, %s, NULL)
            ON CONFLICT (model_id, source_system_id, {}) WHERE modeled_entity_type = {}
            DO UPDATE SET mapping_transformation_document=NULL
            RETURNING mapping_object_id
        """).format(
                    sql.Identifier(f"{layer}_entity_id"),
                    sql.Identifier(f"{layer}_entity_id"),
                    sql.Literal(entity_type),
                ),
                (model_id, entity_type, entity_id, scope.source_system_id),
            ).fetchone()
        )["mapping_object_id"]
        connection.execute(
            sql.SQL("""
            INSERT INTO workflow.mapping_attribute (
                mapping_object_id, model_id, modeled_entity_type, {}, {},
                attribute_mapping_transformation_document
            ) SELECT %s, model_id, %s, {}, {}, '{{}}'::JSONB FROM workflow.{}
              WHERE model_id=%s AND {}=%s
            ON CONFLICT DO NOTHING
        """).format(
                *(
                    sql.Identifier(value)
                    for value in (
                        f"{layer}_entity_id",
                        f"{layer}_attribute_id",
                        f"{layer}_entity_id",
                        f"{layer}_attribute_id",
                        f"{layer}_attribute",
                        f"{layer}_entity_id",
                    )
                )
            ),
            (mapping_id, entity_type, model_id, entity_id),
        )
        connection.execute(
            sql.SQL("""
            INSERT INTO workflow.{} (model_id, {}, {}, {}, {}, {}, {}, {}{})
            VALUES (%s,%s,'postal_code','Postal code','string',2,'active',TRUE{}),
                   (%s,%s,'retired','Retired field','string',3,'inactive',FALSE{})
        """).format(
                *(
                    sql.Identifier(value)
                    for value in (
                        f"{layer}_attribute",
                        f"{layer}_entity_id",
                        f"{layer}_attribute_name",
                        f"{layer}_attribute_definition",
                        f"{layer}_attribute_data_type",
                        f"{layer}_attribute_ordinal_position",
                        f"{layer}_attribute_status",
                        f"{layer}_attribute_is_locked",
                    )
                ),
                sql.SQL(
                    ", dimensional_attribute_role" if layer == "dimensional" else ""
                ),
                sql.SQL(", 'descriptor'" if layer == "dimensional" else ""),
                sql.SQL(", 'descriptor'" if layer == "dimensional" else ""),
            ),
            (model_id, entity_id, model_id, entity_id),
        )
    principal = RequestPrincipal(
        actor_kind=ActorKind.HUMAN,
        entra_tenant_id=actor["entra_tenant_id"],
        entra_object_id=actor["entra_object_id"],
    )
    runtime = WebPostgresDatabase(
        dsn=web_postgres_database.web_runtime_dsn(),
        pool_min=1,
        pool_max=1,
        pool_timeout_seconds=5,
    )
    service = DatabaseMappingReviewService(
        database=runtime,
        authorizer=AuthorizationService(),
        cursor_signing_key=b"disposable-partial-mapping-reads",
    )
    await runtime.open()
    try:
        objects = await service.list_objects(
            principal,
            tenant_id=scope.tenant_id,
            model_id=model_id,
            filters=MappingFilters(entity_type=entity_type),
            page_size=25,
            cursor=None,
        )
        assert [item.mapping_object_id for item in objects.items] == [mapping_id]
        download = await service.export_workbook(
            principal,
            tenant_id=scope.tenant_id,
            model_id=model_id,
            entity_type=entity_type,
            source_system_code=objects.items[0].source_system.system_code,
            expected_model_revision=objects.model_revision,
        )
        workbook = load_workbook(BytesIO(download.content))
        try:
            sheet = workbook.worksheets[0]
            assert sheet["B4"].value == ("Gold" if layer == "dimensional" else "Silver")
            assert sheet["B15"].value == "customer_id"
            assert sheet["B16"].value == "postal_code"
            assert sheet["G16"].value is None
            assert sheet.max_row == 16
        finally:
            workbook.close()
        first = await service.list_attributes(
            principal,
            tenant_id=scope.tenant_id,
            model_id=model_id,
            filters=MappingAttributeFilters(mapping_object_id=mapping_id),
            page_size=1,
            cursor=None,
        )
        assert first.items[0].target.attribute_name == "customer_id"
        assert first.items[0].mapping_attribute_id is not None
        second = await service.list_attributes(
            principal,
            tenant_id=scope.tenant_id,
            model_id=model_id,
            filters=MappingAttributeFilters(mapping_object_id=mapping_id),
            page_size=1,
            cursor=first.next_cursor,
        )
        assert second.next_cursor is None
        missing = second.items[0]
        assert missing.target.attribute_name == "postal_code"
        assert missing.mapping_attribute_id is None
        assert (
            missing.updated_at is None and missing.status is None and missing.is_locked
        )
        locked = await service.list_attributes(
            principal,
            tenant_id=scope.tenant_id,
            model_id=model_id,
            filters=MappingAttributeFilters(mapping_object_id=mapping_id, locked=True),
            page_size=25,
            cursor=None,
        )
        assert locked.items == (missing,)
        with web_postgres_database.connect_owner() as connection:
            connection.execute(
                "UPDATE workflow.mapping_attribute SET "
                "attribute_mapping_transformation_document=NULL WHERE mapping_object_id=%s",
                (mapping_id,),
            )
        empty = await service.list_objects(
            principal,
            tenant_id=scope.tenant_id,
            model_id=model_id,
            filters=MappingFilters(entity_type=entity_type),
            page_size=25,
            cursor=None,
        )
        assert empty.items == ()
        # Existing URLs retain governed history; clearing doesn't erase the record or its identity.
        history = await service.read_object(
            principal,
            tenant_id=scope.tenant_id,
            model_id=model_id,
            mapping_object_id=mapping_id,
        )
        assert history.mapping_document is None
        attribute_history = await service.read_attribute(
            principal,
            tenant_id=scope.tenant_id,
            model_id=model_id,
            mapping_attribute_id=first.items[0].mapping_attribute_id,
        )
        assert attribute_history.mapping_document is None
        targets = await service.list_generation_targets(
            principal,
            tenant_id=scope.tenant_id,
            model_id=model_id,
            entity_type=entity_type,
            page_size=25,
            cursor=None,
        )
        assert any(item.mapping_object_id == mapping_id for item in targets.items)
    finally:
        await runtime.close()
